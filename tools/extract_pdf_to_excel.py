"""
extract_pdf_to_excel.py

Extracts lab tests, prices, durations (result date), and sample collection notes
from the lab price list PDF (قائمة اسعار الفروع (1).pdf) into a clean, structured Excel file.

Rules:
- Truncates / drops any test where price <= 0.
- Extracts duration (days) from 'Result date'.
- Extracts sample collection notes.
- Extracts category/panel headers.
"""

import argparse
import os
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pypdfium2 as pdfium

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass


HEADER_REGEX = re.compile(r'^(.*?)\s+Price\s+Result\s+date\s+Collection\s+notes\s*$', re.I)
TEST_REGEX = re.compile(r'^(.*?)\s+(\d+(?:\.\d+)?)\s*L\.E\.\s*(\d+)\s*Days?(?:\s+(.*))?$', re.I)
BIDI_TEST_REGEX = re.compile(r'^Days\s+(\d+)\s+\.E\.L\s+(\d+(?:\.\d+)?)\s*(.*)$', re.I)


def extract_tests_from_pdf(pdf_path: str, lab_name: str = "معمل الخبراء للتحاليل الطبية"):
    if not Path(pdf_path).exists():
        raise FileNotFoundError(f"PDF file not found: {pdf_path}")

    pdf = pdfium.PdfDocument(pdf_path)
    total_pages = len(pdf)
    print(f"📄 Processing PDF: {pdf_path} ({total_pages} pages)")

    valid_tests = []
    truncated_zero_price = []
    current_category = "General"

    for p in range(total_pages):
        text = pdf[p].get_textpage().get_text_range()
        lines = [l.strip() for l in text.splitlines() if l.strip()]

        for line in lines:
            # Skip page headers / footers
            if "Printed In" in line or "قائمة أسعار" in line or "Page Current" in line or ("from" in line and "Page" in line):
                continue
            if "@" in line and "mail" in line:
                continue

            # Check for Category Header
            hm = HEADER_REGEX.match(line)
            if hm:
                current_category = hm.group(1).strip()
                continue

            # Standard Test Row
            tm = TEST_REGEX.match(line)
            if tm:
                name, price_str, duration_str, notes = tm.groups()
                price = float(price_str)
                duration = int(duration_str)
                notes_clean = (notes or "").strip()

                if price <= 0:
                    truncated_zero_price.append({
                        "category": current_category,
                        "name": name.strip(),
                        "price": price,
                        "duration": duration,
                        "notes": notes_clean,
                        "page": p + 1
                    })
                else:
                    valid_tests.append({
                        "lab_name": lab_name,
                        "category": current_category,
                        "name": name.strip(),
                        "price": price,
                        "duration_days": duration,
                        "sample_notes": notes_clean,
                    })
                continue

            # Bidi Arabic Test Row (e.g. Days 0 .E.L 50 زيارة منزلية)
            bm = BIDI_TEST_REGEX.match(line)
            if bm:
                duration_str, price_str, name = bm.groups()
                price = float(price_str)
                duration = int(duration_str)

                if price <= 0:
                    truncated_zero_price.append({
                        "category": current_category,
                        "name": name.strip(),
                        "price": price,
                        "duration": duration,
                        "notes": "",
                        "page": p + 1
                    })
                else:
                    valid_tests.append({
                        "lab_name": lab_name,
                        "category": current_category,
                        "name": name.strip(),
                        "price": price,
                        "duration_days": duration,
                        "sample_notes": "",
                    })

    return valid_tests, truncated_zero_price


def save_to_excel(tests: list[dict], zero_price_tests: list[dict], output_excel_path: str):
    wb = openpyxl.Workbook()

    # --- Sheet 1: Valid Tests ---
    ws = wb.active
    ws.title = "Lab Tests"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "Lab Name",
        "Category",
        "Test Name",
        "Price (L.E.)",
        "Duration (Days)",
        "Sample Type & Collection Notes"
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    ws.append(headers)
    ws.row_dimensions[1].height = 26

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for row_idx, test in enumerate(tests, start=2):
        row_values = [
            test["lab_name"],
            test["category"],
            test["name"],
            test["price"],
            test["duration_days"],
            test["sample_notes"],
        ]
        ws.append(row_values)
        ws.row_dimensions[row_idx].height = 20

        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, len(row_values) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill

            if col_idx in (1, 2, 3):
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 4:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00"
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = "0"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Column widths
    col_widths = {1: 28, 2: 24, 3: 40, 4: 16, 5: 16, 6: 45}
    for col_idx, width in col_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # --- Sheet 2: Truncated Zero-Price Tests Log ---
    ws_zero = wb.create_sheet(title="Truncated Zero-Price Tests")
    ws_zero.views.sheetView[0].showGridLines = True
    zero_headers = ["Page", "Category", "Test Name", "Price", "Duration", "Notes", "Action Taken"]
    ws_zero.append(zero_headers)
    ws_zero.row_dimensions[1].height = 24

    for col_idx in range(1, len(zero_headers) + 1):
        cell = ws_zero.cell(row=1, column=col_idx)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for z_idx, z in enumerate(zero_price_tests, start=2):
        ws_zero.append([
            z["page"],
            z["category"],
            z["name"],
            z["price"],
            z["duration"],
            z["notes"],
            "Truncated / Dropped (Price is 0)"
        ])
        for c in range(1, 8):
            ws_zero.cell(row=z_idx, column=c).border = thin_border

    zero_widths = {1: 10, 2: 24, 3: 35, 4: 12, 5: 12, 6: 35, 7: 30}
    for col_idx, width in zero_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws_zero.column_dimensions[col_letter].width = width

    wb.save(output_excel_path)
    print(f"✅ Excel file saved successfully: {output_excel_path}")


def main():
    parser = argparse.ArgumentParser(description="Extract lab tests from PDF to Excel")
    parser.add_argument("--pdf", default="قائمة اسعار الفروع (1).pdf", help="Source PDF file path")
    parser.add_argument("--out", default="extracted_lab_tests.xlsx", help="Target output Excel file path")
    parser.add_argument("--lab-name", default="معمل الخبراء للتحاليل الطبية", help="Name of laboratory")
    args = parser.parse_args()

    tests, zero_price = extract_tests_from_pdf(args.pdf, lab_name=args.lab_name)
    print(f"📊 Extracted {len(tests)} valid tests (with price > 0)")
    print(f"✂️  Truncated {len(zero_price)} tests with price <= 0")

    save_to_excel(tests, zero_price, args.out)


if __name__ == "__main__":
    main()
