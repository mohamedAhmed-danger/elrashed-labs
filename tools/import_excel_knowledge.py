"""
bulk_import_tests / import_excel_knowledge.py

Script لعمل استيراد وتوليد لمعرفة التحاليل من ملف Excel (مثل extracted_lab_tests.xlsx)
باستخدام نفس الـ AI pipeline:
    generate_test()  ->  TestsService.create_lab_service()

المميزات:
1. قراءة الملف الجديد المستخرج من PDF (6 أعمدة: Lab Name, Category, Test Name, Price, Duration, Notes)
   ودعم الملفات السابقة (3 أعمدة: Type, Name, Price).
2. استبعاد / ترنكيت أي تحليل سعره صفر (Price <= 0).
3. أخذ مدة التحليل (Duration) من الإكسيل مباشرة، وتحويلها لساعات تلقائيًا (مثل: 1 يوم = 24 ساعة، 0 يوم = 12 ساعة)
   لتتوافق مع حقل duration في الداتابيز وواجهة لوحة التحكم.
4. إرسال ملاحظات ونوع العينة (Sample / Collection Notes) لـ Gemini أثناء التوليد لتحسين دقة التعليمات.
5. تحديد المعمل المستهدف عبر --lab-id أو --lab-name أو أول معمل افتراضيًا.
6. Resume تلقائي، تسجيل الفشل في CSV، تسجيل تفصيلي في TXT، وخيارات --dry-run و --skip-vectors.
"""

import argparse
import csv
import sys
import time
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import openpyxl

try:
    try:
        from app import app
    except ImportError:
        from app import create_app
        app = create_app()

    from models.models import db, Laboratory, LabService
    import services.tests_service as tests_service_module
    from services.domain.tests_service import TestsService
    from services.shared.generation_service import generate_test
    from services.shared.vector_service import upsert_test_vector as _real_upsert_test_vector
except ImportError:
    app = None
    db = None
    Laboratory = None
    LabService = None
    tests_service_module = None
    TestsService = None
    generate_test = None
    _real_upsert_test_vector = None

try:
    from utils.text_utils import build_test_text
except ImportError:
    def build_test_text(name, description, keywords):
        return f"{name}. {description or ''} {' '.join(keywords or [])}".strip()


# --------------------------------------------------------------------------
# إعدادات افتراضية
# --------------------------------------------------------------------------
DEFAULT_EXCEL_PATH = "extracted_lab_tests.xlsx"
FAILED_LOG_PATH = "failed_tests.csv"
DEFAULT_DETAIL_LOG_PATH = "import_log.txt"
SLEEP_BETWEEN_CALLS = 1.0


def get_target_laboratory(lab_id: int | None = None, lab_name: str | None = None) -> Laboratory:
    """
    جلب المعمل المستهدف بناءً على ID أو الاسم، أو المعمل الأول كخيار افتراضي.
    """
    if lab_id:
        lab = db.session.get(Laboratory, lab_id)
        if not lab:
            labs = Laboratory.query.all()
            available = ", ".join(f"[{l.id}: {l.name}]" for l in labs) or "لا يوجد معامل"
            raise ValueError(f"❌ لم يتم العثور على معمل بالمعرف {lab_id}. المعامل المتاحة: {available}")
        return lab

    if lab_name:
        lab = Laboratory.query.filter(Laboratory.name.ilike(f"%{lab_name.strip()}%")).first()
        if not lab:
            labs = Laboratory.query.all()
            available = ", ".join(f"[{l.id}: {l.name}]" for l in labs) or "لا يوجد معامل"
            raise ValueError(f"❌ لم يتم العثور على معمل بالاسم '{lab_name}'. المعامل المتاحة: {available}")
        return lab

    lab = Laboratory.query.order_by(Laboratory.id.asc()).first()
    if not lab:
        raise RuntimeError(
            "❌ لا يوجد أي معمل (Laboratory) في قاعدة البيانات.\n"
            "يرجى إضافة معمل واحد على الأقل قبل تشغيل الاسكريبت."
        )
    return lab


def read_excel_rows(excel_path: str):
    """
    قراءة ملف الإكسيل مع استخراج:
    - الاسم (name)
    - السعر (price) - واستبعاد السعر صفر
    - المدة (duration_days)
    - نوع العينة والملاحظات (sample_notes)
    - التصنيف (category)
    - اسم المعمل (lab_name)
    """
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # اختيار الشيت المناسب
    if "Lab Tests" in wb.sheetnames:
        ws = wb["Lab Tests"]
    else:
        ws = wb.active

    rows = []
    zero_price_count = 0

    header = [str(cell).strip().lower() for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    is_6_cols = len(header) >= 6 or any("duration" in h for h in header)

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(row):
            continue

        if is_6_cols:
            lab_name_val = row[0]
            category_val = row[1]
            name_val = row[2]
            price_val = row[3]
            duration_val = row[4]
            notes_val = row[5] if len(row) > 5 else ""

            if not name_val or str(name_val).strip() == "":
                continue

            try:
                price = float(price_val) if price_val not in (None, "") else 0.0
            except (ValueError, TypeError):
                price = 0.0

            # استبعاد وتجاهل التحاليل بسعر صفر (Truncate zero-price)
            if price <= 0.0:
                zero_price_count += 1
                continue

            try:
                duration_days = int(duration_val) if duration_val not in (None, "") else None
            except (ValueError, TypeError):
                duration_days = None

            rows.append({
                "lab_name": str(lab_name_val).strip() if lab_name_val else "",
                "category": str(category_val).strip() if category_val else "General",
                "name": str(name_val).strip(),
                "price": price,
                "duration_days": duration_days,
                "sample_notes": str(notes_val).strip() if notes_val else "",
            })
        else:
            # النسخة السابقة (3 أعمدة: Type, Name, Price)
            row_type, name_val, price_val = row[0], row[1], row[2]
            if not name_val or str(name_val).strip() == "":
                continue

            try:
                price = float(price_val) if price_val not in (None, "") else 0.0
            except (ValueError, TypeError):
                price = 0.0

            if price <= 0.0:
                zero_price_count += 1
                continue

            rows.append({
                "lab_name": "",
                "category": str(row_type).strip() if row_type else "General",
                "name": str(name_val).strip(),
                "price": price,
                "duration_days": None,
                "sample_notes": "",
            })

    if zero_price_count > 0:
        print(f"✂️  تم استبعاد {zero_price_count} تحليلاً تلقائيًا لأن سعرها 0.")

    return rows


def already_exists(name: str, laboratory_id: int | None = None) -> bool:
    """التحقق مما إذا كان التحليل مسجلاً بالفعل لتجنب التكرار واستهلاك الـ AI."""
    q = LabService.query.filter_by(name=name)
    if laboratory_id:
        q = q.filter_by(laboratory_id=laboratory_id)
    return q.first() is not None


def log_failure(name: str, price: float, error: str):
    file_exists = Path(FAILED_LOG_PATH).exists()
    with open(FAILED_LOG_PATH, "a", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        if not file_exists:
            writer.writerow(["name", "price", "error"])
        writer.writerow([name, price, error])


def generation_result_to_dict(generated) -> dict:
    if hasattr(generated, "model_dump"):
        return generated.model_dump()
    return generated.dict()


def format_list(value) -> str:
    if not value:
        return "[]"
    return "[" + ", ".join(str(v) for v in value) + "]"


def write_detail_log(log_path: str, index: int, total: int, name: str,
                      generated_dict: dict | None,
                      embed_text: str | None,
                      db_fields: dict | None,
                      db_status: str,
                      qdrant_payload: dict | None,
                      qdrant_status: str):
    lines = [
        "=" * 70,
        f"[{index}/{total}] {name}",
        "=" * 70,
        "🧬 نتيجة التوليد من Gemini:"
    ]
    if generated_dict is None:
        lines.append("  (لم يتم التوليد - حصل خطأ قبل الوصول للخطوة دي)")
    else:
        lines.append(f"  description          : {generated_dict.get('description', '')}")
        lines.append(f"  patient_instructions  : {generated_dict.get('patient_instructions', '')}")
        lines.append(f"  keywords              : {format_list(generated_dict.get('keywords'))}")
        lines.append(f"  alias_name            : {format_list(generated_dict.get('alias_name'))}")
        lines.append(f"  duration              : {generated_dict.get('duration')}")
        lines.append(f"  sample_type           : {generated_dict.get('sample_type')}")
    lines.append("")

    lines.append("🧠 النص الذي تم تحويله لـ Embedding:")
    lines.append(f'  "{embed_text}"' if embed_text else "  (لم يُحسب)")
    lines.append("")

    lines.append("💾 الحفظ في الداتابيز (LabService):")
    if db_fields is None:
        lines.append("  (لم يُحفظ)")
    else:
        for key, val in db_fields.items():
            if isinstance(val, list):
                val = format_list(val)
            lines.append(f"  {key:<20}: {val}")
    lines.append(f"  الحالة: {db_status}")
    lines.append("")

    lines.append("🔎 بيانات الفيكتور لـ Qdrant:")
    if qdrant_payload is None:
        lines.append("  (لم يُرسل)")
    else:
        for key, val in qdrant_payload.items():
            if isinstance(val, list):
                val = format_list(val)
            lines.append(f"  {key:<12}: {val}")
    lines.append(f"  الحالة: {qdrant_status}")
    lines.append("\n")

    with open(log_path, "a", encoding="utf-8-sig") as f:
        f.write("\n".join(lines))


def main():
    parser = argparse.ArgumentParser(description="Bulk import lab tests from Excel using the AI generation pipeline")
    parser.add_argument("--excel", default=DEFAULT_EXCEL_PATH, help=f"Path to Excel file (default: {DEFAULT_EXCEL_PATH})")
    parser.add_argument("--lab-id", type=int, default=None, help="Target Laboratory ID in the database")
    parser.add_argument("--lab-name", type=str, default=None, help="Target Laboratory Name in the database")
    parser.add_argument("--limit", type=int, default=None, help="Only process the first N rows (for testing)")
    parser.add_argument("--sleep", type=float, default=SLEEP_BETWEEN_CALLS, help="Seconds to sleep between AI calls")
    parser.add_argument("--duration-unit", choices=["hours", "days"], default="hours",
                        help="Convert duration to hours (e.g. 1 day = 24 hrs) or store raw days (default: hours)")
    parser.add_argument("--dry-run", action="store_true", help="Generate with AI but don't insert into DB or Qdrant")
    parser.add_argument("--skip-vectors", action="store_true", help="Insert into DB only without connecting to Qdrant")
    parser.add_argument("--log-file", default=DEFAULT_DETAIL_LOG_PATH, help="Path to detailed per-test log file")
    args = parser.parse_args()

    if not Path(args.excel).exists():
        print(f"❌ الملف غير موجود: {args.excel}")
        sys.exit(1)

    last_qdrant_call = {}

    def logged_upsert_test_vector(test_id, test_name, description, keywords):
        try:
            embed_text = build_test_text(test_name, description, keywords)
        except Exception as e:
            embed_text = f"<تعذر بناء نص الـ embedding: {e}>"

        payload = {
            "test_id": test_id,
            "name": test_name,
            "description": description,
            "keywords": keywords or [],
        }

        if args.skip_vectors:
            last_qdrant_call.update({
                "embed_text": embed_text,
                "payload": payload,
                "status": "تخطي (--skip-vectors)"
            })
            return

        try:
            _real_upsert_test_vector(test_id, test_name, description, keywords)
            last_qdrant_call.update({
                "embed_text": embed_text,
                "payload": payload,
                "status": "✅ تم الـ upsert في Qdrant بنجاح"
            })
        except Exception as e:
            last_qdrant_call.update({
                "embed_text": embed_text,
                "payload": payload,
                "status": f"❌ فشل Qdrant: {e}"
            })

    # استبدال الـ upsert بالـ wrapper المسجل
    tests_service_module.upsert_test_vector = logged_upsert_test_vector

    if args.skip_vectors:
        print("⚠️  --skip-vectors مفعّل: لن يتم عمل upsert في Qdrant (سيتم الحفظ في PostgreSQL فقط).")

    with app.app_context():
        try:
            lab = get_target_laboratory(lab_id=args.lab_id, lab_name=args.lab_name)
        except Exception as e:
            print(e)
            sys.exit(1)

        print(f"➡️  المعمل المستهدف: {lab.name} (id={lab.id})")

        rows = read_excel_rows(args.excel)
        if args.limit:
            rows = rows[: args.limit]

        total = len(rows)
        print(f"📄 عدد التحاليل الجاهزة للمعالجة: {total}")

        created, skipped, failed = 0, 0, 0

        try:
            for i, row in enumerate(rows, start=1):
                name = row["name"]
                price = row["price"]
                raw_duration_days = row.get("duration_days")
                sample_notes = row.get("sample_notes", "")

                print(f"[{i}/{total}] {name} (السعر: {price} ج.م) ...", end=" ", flush=True)

                if already_exists(name, laboratory_id=lab.id):
                    print("⏭️  موجود بالفعل - تخطي")
                    skipped += 1
                    continue

                generated_dict = None
                db_fields = None
                db_status = "(قيد المعالجة)"
                last_qdrant_call.clear()

                try:
                    # استدعاء Gemini مع تزويده بملاحظات العينة إن وجدت لدقة أكبر
                    generated = generate_test(
                        name=name,
                        instructions=sample_notes if sample_notes else None
                    )
                    generated_dict = generation_result_to_dict(generated)

                    # تحديد مدة النتيجة (بالساعات أو الأيام)
                    if raw_duration_days is not None:
                        if args.duration_unit == "hours":
                            # 1 يوم = 24 ساعة، 0 يوم = 12 ساعة كحد أدنى لنفس اليوم
                            calculated_duration = raw_duration_days * 24 if raw_duration_days > 0 else 12
                        else:
                            calculated_duration = raw_duration_days
                    else:
                        calculated_duration = getattr(generated, "duration", 24) or 24

                    # تحديد نوع العينة (إذا تم استخراجه من الملف أو من الـ AI)
                    final_sample_type = (generated.sample_type or "").strip()
                    if not final_sample_type and sample_notes:
                        final_sample_type = sample_notes

                    db_fields = {
                        "laboratory_id": lab.id,
                        "name": name,
                        "price": price,
                        "description": generated.description,
                        "patient_instructions": generated.patient_instructions,
                        "duration": calculated_duration,
                        "sample_type": final_sample_type,
                        "keywords": generated.keywords,
                        "alias_name": generated.alias_name,
                    }

                    if args.dry_run:
                        db_status = "(dry-run - لم يُحفظ في الداتابيز)"
                        embed_text = build_test_text(name, generated.description, generated.keywords)
                        last_qdrant_call.update({
                            "embed_text": embed_text,
                            "payload": {
                                "test_id": None,
                                "name": name,
                                "description": generated.description,
                                "keywords": generated.keywords or [],
                            },
                            "status": "(dry-run)",
                        })
                        print("✅ (dry-run نجح التوليد)")
                        created += 1
                        write_detail_log(
                            args.log_file, i, total, name, generated_dict,
                            last_qdrant_call.get("embed_text"), db_fields, db_status,
                            last_qdrant_call.get("payload"), last_qdrant_call.get("status"),
                        )
                        time.sleep(args.sleep)
                        continue

                    # الحفظ الفعلي عبر TestsService
                    tests_service = TestsService(laboratory_id=lab.id)
                    new_test, message = tests_service.create_lab_service(
                        laboratory_id=lab.id,
                        name=name,
                        price=price,
                        description=generated.description,
                        patient_instructions=generated.patient_instructions,
                        duration=calculated_duration,
                        sample_type=final_sample_type,
                        keywords=generated.keywords,
                        alias_name=generated.alias_name,
                    )

                    if new_test:
                        print("✅ تم الحفظ")
                        db_status = "✅ تم الحفظ بنجاح"
                        created += 1
                    else:
                        print(f"⚠️ فشل: {message}")
                        db_status = f"❌ فشل: {message}"
                        log_failure(name, price, message)
                        failed += 1

                    write_detail_log(
                        args.log_file, i, total, name, generated_dict,
                        last_qdrant_call.get("embed_text"), db_fields, db_status,
                        last_qdrant_call.get("payload"), last_qdrant_call.get("status", "(لم يحدث)"),
                    )

                except KeyboardInterrupt:
                    raise
                except Exception as e:
                    print(f"❌ خطأ: {e}")
                    log_failure(name, price, str(e))
                    write_detail_log(
                        args.log_file, i, total, name, generated_dict,
                        last_qdrant_call.get("embed_text"), db_fields, f"❌ خطأ: {e}",
                        last_qdrant_call.get("payload"), last_qdrant_call.get("status", "(لم يحدث)"),
                    )
                    failed += 1

                time.sleep(args.sleep)

        except KeyboardInterrupt:
            print("\n\n⏹️ تم إيقاف الاسكريبت يدويًا (Ctrl+C).")
            print("ما تم حفظه مسجل بالفعل في الداتابيز، وعند إعادة التشغيل سيستكمل من حيث توقف.")

        print("\n----------------------------------------")
        print(f"✅ تم الإدراج: {created}")
        print(f"⏭️ تم التخطي: {skipped}")
        print(f"❌ فشل:       {failed}")
        if failed:
            print(f"📄 مراجعة التحاليل الفاشلة في: {FAILED_LOG_PATH}")
        print(f"📝 السجل التفصيلي للتوليد في: {args.log_file}")
        print("----------------------------------------")


if __name__ == "__main__":
    main()